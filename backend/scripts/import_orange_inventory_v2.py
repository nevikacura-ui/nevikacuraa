"""
Import Orange Pharmacy inventory from XLSX + CSV files.
Wipes old medicines + pharmacy_inventory, inserts fresh filtered data.
Run: python3 scripts/import_orange_inventory_v2.py
"""
import asyncio
import csv
import os
import re
import uuid
from datetime import datetime, timezone
from motor.motor_asyncio import AsyncIOMotorClient

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl"])
    import openpyxl

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")

XLSX_PATH = "/app/backend/data/imports/imp_new.xlsx"
CSV_PATH = "/tmp/onemg.csv"

# Exclusion keywords (case-insensitive)
EXCLUDE_PATTERNS = re.compile(r'\b(inj|injection|injections|consultation)\b', re.IGNORECASE)

def parse_price(val):
    """Parse price from string or number, stripping currency symbols."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    cleaned = re.sub(r'[₹?\s,]', '', str(val).strip())
    try:
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0

def apply_discount(mrp, pct=17.5):
    """Apply discount and round to 2 decimals."""
    return round(mrp * (1 - pct / 100), 2)

def should_exclude(name):
    """Check if item should be excluded based on keywords."""
    return bool(EXCLUDE_PATTERNS.search(name or ''))

def load_xlsx_items():
    """Load items from the Orange Pharmacy XLSX stock report."""
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    
    # Headers are at row 13
    headers = [str(cell.value or '').strip() for cell in ws[13]]
    print(f"XLSX headers: {headers[:15]}")
    
    items = []
    seen = set()
    
    for row in ws.iter_rows(min_row=14, values_only=True):
        name = str(row[0] or '').strip()
        if not name:
            continue
        
        # Skip excluded items
        if should_exclude(name):
            continue
        
        # Deduplicate by name (case-insensitive)
        name_key = name.lower()
        if name_key in seen:
            continue
        seen.add(name_key)
        
        form = str(row[1] or '').strip()
        company = str(row[2] or '').strip()
        barcode = str(row[3] or '').strip() if row[3] else ''
        mrp = parse_price(row[11])  # MRP (Rs) column
        
        if mrp <= 0:
            mrp = parse_price(row[10])  # fallback to Old MRP
        
        sale_price = apply_discount(mrp)
        discount_pct = 17.5
        
        items.append({
            "name": name,
            "form": form,
            "company": company,
            "barcode": barcode,
            "mrp": mrp,
            "sale_price": sale_price,
            "discount_percent": discount_pct,
            "source": "orange_pharmacy_stock",
        })
    
    print(f"XLSX: Loaded {len(items)} items (after filtering + dedup)")
    return items, seen

def load_csv_items(existing_names):
    """Load items from the 1mg CSV, skipping duplicates already in XLSX."""
    items = []
    
    with open(CSV_PATH, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = (row.get('Drug_Name') or '').strip()
            if not name:
                continue
            
            # Skip excluded items
            if should_exclude(name):
                continue
            
            # Skip if already in XLSX data
            name_key = name.lower()
            if name_key in existing_names:
                continue
            existing_names.add(name_key)
            
            drug_type = (row.get('Drug_Type') or '').strip()
            manufacturer = (row.get('Manufacturer') or row.get('Marketer') or '').strip()
            mrp = parse_price(row.get('MRP', '0'))
            
            if mrp <= 0:
                continue
            
            sale_price = apply_discount(mrp)
            discount_pct = 17.5
            
            items.append({
                "name": name,
                "form": drug_type,
                "company": manufacturer,
                "barcode": "",
                "mrp": mrp,
                "sale_price": sale_price,
                "discount_percent": discount_pct,
                "source": "1mg_catalog",
            })
    
    print(f"CSV: Loaded {len(items)} additional items (after filtering + dedup)")
    return items

async def main():
    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]
    
    print("=" * 60)
    print("ORANGE PHARMACY INVENTORY IMPORT v2")
    print("=" * 60)
    
    # Step 1: Wipe old collections
    old_medicines = await db.medicines.count_documents({})
    old_inventory = await db.pharmacy_inventory.count_documents({})
    print(f"\nOld data: medicines={old_medicines}, pharmacy_inventory={old_inventory}")
    
    if old_medicines > 0:
        await db.medicines.drop()
        print(f"  Dropped 'medicines' collection ({old_medicines} docs)")
    
    if old_inventory > 0:
        await db.pharmacy_inventory.drop()
        print(f"  Dropped 'pharmacy_inventory' collection ({old_inventory} docs)")
    
    # Step 2: Load from XLSX
    xlsx_items, seen_names = load_xlsx_items()
    
    # Step 3: Load from CSV (skipping dupes)
    csv_items = load_csv_items(seen_names)
    
    all_items = xlsx_items + csv_items
    print(f"\nTotal combined items: {len(all_items)}")
    
    # Step 4: Prepare MongoDB documents
    now = datetime.now(timezone.utc).isoformat()
    docs = []
    for item in all_items:
        docs.append({
            "id": str(uuid.uuid4()),
            "name": item["name"],
            "form": item["form"],
            "company": item["company"],
            "category": "",
            "image_url": "",
            "barcode": item["barcode"],
            "mrp": item["mrp"],
            "sale_price": item["sale_price"],
            "discount_percent": item["discount_percent"],
            "stock": 100,
            "unit": item["form"] or "Unit",
            "description": "",
            "source": item["source"],
            "created_at": now,
            "created_by": "inventory_import_v2",
        })
    
    # Step 5: Bulk insert
    if docs:
        result = await db.pharmacy_inventory.insert_many(docs)
        print(f"\nInserted {len(result.inserted_ids)} items into pharmacy_inventory")
    
    # Step 6: Create index for fast search
    await db.pharmacy_inventory.create_index([("name", 1)])
    await db.pharmacy_inventory.create_index([("company", 1)])
    await db.pharmacy_inventory.create_index([("category", 1)])
    
    # Step 7: Verify
    final_count = await db.pharmacy_inventory.count_documents({})
    print(f"\nFinal pharmacy_inventory count: {final_count}")
    
    # Sample output
    print("\nSample items:")
    async for doc in db.pharmacy_inventory.find({}, {"_id": 0, "name": 1, "mrp": 1, "sale_price": 1, "source": 1}).limit(8):
        print(f"  {doc['name']}: MRP ₹{doc['mrp']} → ₹{doc['sale_price']} ({doc['source']})")
    
    # Source breakdown
    xlsx_count = await db.pharmacy_inventory.count_documents({"source": "orange_pharmacy_stock"})
    csv_count = await db.pharmacy_inventory.count_documents({"source": "1mg_catalog"})
    print(f"\nBreakdown: XLSX={xlsx_count}, CSV={csv_count}")
    
    client.close()
    print("\nDone!")

if __name__ == "__main__":
    asyncio.run(main())
