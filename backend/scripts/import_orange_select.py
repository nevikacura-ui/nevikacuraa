"""
Import Orange Select inventory from 3 sources:
1. imp_new.xlsx - Orange Pharmacy detailed stock (562 unique products with full details)
2. current_stock_new.xlsx - DiaGyn Healthcare stock (216 unique products with full details)
3. stock_summary_new.pdf - Orange Pharmacy stock summary (858 items, name + sale_price only)

Pricing rules:
- DiaGyn (Local Well): MRP is the base rate
- Vyapaar (Orange Pharmacy): sale_rate + 15% = orange MRP
- Duplicates: prefer DiaGyn (Local Well)

Total expected: ~1500+ unique items
"""

import asyncio
import os
import sys
import re
import uuid
from datetime import datetime, timezone
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), '..', '.env'))

import openpyxl

try:
    import fitz  # PyMuPDF
except ImportError:
    print("Installing pymupdf...")
    os.system("pip install pymupdf -q")
    import fitz

from motor.motor_asyncio import AsyncIOMotorClient

MONGO_URL = os.environ.get('MONGO_URL')
DB_NAME = os.environ.get('DB_NAME')
DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'data', 'imports')

# Category inference based on medicine names and compositions
CATEGORY_MAP = {
    'diabetes': ['glimep', 'metformin', 'vogli', 'dailyglim', 'glycomet', 'jalra', 'teneli', 'glimy', 'insulin', 'diabet'],
    'heart & bp care': ['telmi', 'amlo', 'aten', 'losart', 'olmesar', 'corbis', 'telmik', 'cardio', 'ecosprin', 'cilni', 'metropol'],
    'respiratory care': ['montek', 'budamate', 'aerocort', 'duolin', 'seroflo', 'foracort', 'levolin', 'salbut', 'asthma', 'bronch'],
    'gastro & digestive': ['pantop', 'omee', 'rablet', 'acigene', 'freego', 'gutbliss', 'livoluk', 'visco', 'antacid', 'enzym', 'zymvax'],
    'antibiotics': ['amoxicillin', 'azithro', 'cefixime', 'cipro', 'oflox', 'doxo', 'augment', 'omnicl', 'monocef', 'aristomox', 'brutacross', 'aliclair', 'zoclar', 'zithowin'],
    'vitamins & supplements': ['vitamin', 'calcium', 'calcigen', 'folic', 'iron', 'fercee', 'fericip', 'multirich', 'megavax', 'curivit', 'threptin', 'nutreme', 'protinules'],
    'pain relief': ['dolo', 'paracet', 'ibuprofen', 'zerodol', 'napro', 'adigesic', 'dolobreak', 'dolozox', 'etozox', 'flexon', 'combi'],
    'skin & dermatology': ['cream', 'oint', 'lotion', 'calamine', 'acne', 'fungitop', 'mupicip', 'melnor', 'burnheal', 'stayclear', 'ring guard', 'krimly', 'calen'],
    'allergy & cold': ['cetriz', 'okacet', 'l hist', 'allegra', 'allerfex', 'cheston', 'alkof', 'kofclear', 'cuflift', 'zukamin', 'koldkind', 'anergen'],
    'thyroid care': ['thyro', 'thiro'],
    'women\'s health': ['femipris', 'pregvom', 'primepill', 'myo bd', 'myovil', 'funtwist', 'caldimint fem'],
    'bones & joints': ['tendon', 'chymotom', 'chymori', 'nafodil'],
    'anti-fungal': ['fungicip', 'fungiforce', 'terbizol', 'ring out', 'candid', 'ketocip', 'ketonalog'],
    'neuro & brain': ['gabakind', 'pregaban', 'pregnerv', 'pre gabaprex', 'vertiford', 'vertigil', 'vertiron'],
    'eye & ear care': ['eyedrop', 'optibes', 'moxigram'],
    'urology & kidney': ['roliten', 'soliwise', 'pyridium'],
    'liver care': ['hepaco', 'livosoft', 'silymarin'],
    'hair care': ['hairfolic', 'hairfollic', '8x kt', 'head&shoulder', 'ketocip shampoo'],
    'personal care': ['colgate', 'veet', 'sofy', 'whisper', 'venus', 'liveasy', 'eveready', 'olivo', 'happy baby', 'dabur'],
    'probiotics': ['rinilact', 'sporlac', 'gutpro'],
}

def infer_category(name, composition='', form=''):
    """Infer category based on product name, composition, and form"""
    text = f"{name} {composition} {form}".lower()
    for category, keywords in CATEGORY_MAP.items():
        for kw in keywords:
            if kw in text:
                return category.title()
    # Default based on form
    form_lower = (form or '').lower()
    if any(f in form_lower for f in ['cream', 'oint', 'lotion', 'gel']):
        return 'Skin & Dermatology'
    if any(f in form_lower for f in ['shampoo']):
        return 'Hair Care'
    if any(f in form_lower for f in ['drop']):
        return 'Eye & Ear Care'
    return 'General Medicine'

def infer_form(name):
    """Infer dosage form from product name"""
    name_lower = name.lower()
    form_map = [
        ('tab', 'Tablet'), ('cap', 'Capsule'), ('syp', 'Syrup'), ('syrp', 'Syrup'),
        ('syrup', 'Syrup'), ('cream', 'Cream'), ('oint', 'Ointment'), ('gel', 'Gel'),
        ('drop', 'Drops'), ('drp', 'Drops'), ('inj', 'Injection'), ('susp', 'Suspension'),
        ('powder', 'Powder'), ('pwdr', 'Powder'), ('sachet', 'Sachet'), ('lotion', 'Lotion'),
        ('shampoo', 'Shampoo'), ('spray', 'Spray'), ('liquid', 'Liquid'),
        ('facewash', 'Facewash'), ('soap', 'Soap'), ('sup', 'Suppository'),
        ('kit', 'Kit'), ('granule', 'Granules'),
    ]
    for keyword, form in form_map:
        if keyword in name_lower:
            return form
    return ''

def parse_orange_xlsx(filepath):
    """Parse Orange Pharmacy XLSX with full product details"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Sheet 1']
    products = {}
    
    for row in ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True):
        name = row[0]
        if not name or not isinstance(name, str) or not name.strip():
            continue
        
        name = name.strip()
        key = name.lower()
        
        mrp = row[11] or 0
        sale_rate = row[12] or mrp
        net_rate = row[15] or 0
        
        try:
            mrp = float(mrp)
        except (ValueError, TypeError):
            mrp = 0
        try:
            sale_rate = float(sale_rate)
        except (ValueError, TypeError):
            sale_rate = mrp
        try:
            net_rate = float(net_rate)
        except (ValueError, TypeError):
            net_rate = 0
        
        # Orange pricing: sale_rate + 15% = orange MRP
        orange_price = round(sale_rate * 1.15, 2) if sale_rate > 0 else round(mrp * 1.15, 2)
        
        units = row[17]
        try:
            units = int(units) if units else 1
        except (ValueError, TypeError):
            units = 1
        
        if key not in products or mrp > (products[key].get('mrp', 0) or 0):
            form = row[1] or ''
            company = row[2] or ''
            composition = row[20] or ''
            category = row[6] or ''
            rx_type = row[4] or ''
            
            if not category:
                category = infer_category(name, composition, form)
            
            products[key] = {
                'name': name,
                'form': form.strip() if form else '',
                'company': company.strip() if company else '',
                'rx_type': rx_type.strip() if rx_type else '',
                'mrp': mrp,
                'sale_price': sale_rate,
                'diagyn_price': mrp,  # DiaGyn uses MRP
                'orange_price': orange_price,
                'net_rate': net_rate,
                'units_per_pack': units,
                'composition': composition.strip() if composition else '',
                'generic_name': '',
                'category': category.strip() if category else 'General Medicine',
                'store': 'orange_pharmacy',
                'formulary_source': 'vyapaar',
            }
    
    return products

def parse_diagyn_xlsx(filepath):
    """Parse DiaGyn Healthcare XLSX with full product details"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb['Sheet 1']
    products = {}
    
    for row in ws.iter_rows(min_row=14, max_row=ws.max_row, values_only=True):
        name = row[0]
        if not name or not isinstance(name, str) or not name.strip():
            continue
        
        name = name.strip()
        key = name.lower()
        
        mrp = row[11] or 0
        sale_rate = row[12] or mrp
        net_rate = row[15] or 0
        
        try:
            mrp = float(mrp)
        except (ValueError, TypeError):
            mrp = 0
        try:
            sale_rate = float(sale_rate)
        except (ValueError, TypeError):
            sale_rate = mrp
        try:
            net_rate = float(net_rate)
        except (ValueError, TypeError):
            net_rate = 0
        
        units = row[17]
        try:
            units = int(units) if units else 1
        except (ValueError, TypeError):
            units = 1
        
        if key not in products or mrp > (products[key].get('mrp', 0) or 0):
            form = row[1] or ''
            company = row[2] or ''
            composition = row[20] or ''
            category = row[6] or ''
            rx_type = row[4] or ''
            
            if not category:
                category = infer_category(name, composition, form)
            
            products[key] = {
                'name': name,
                'form': form.strip() if form else '',
                'company': company.strip() if company else '',
                'rx_type': rx_type.strip() if rx_type else '',
                'mrp': mrp,
                'sale_price': mrp,  # DiaGyn uses MRP as sale price
                'diagyn_price': mrp,
                'orange_price': round(sale_rate * 1.15, 2) if sale_rate > 0 else round(mrp * 1.15, 2),
                'net_rate': net_rate,
                'units_per_pack': units,
                'composition': composition.strip() if composition else '',
                'generic_name': '',
                'category': category.strip() if category else 'General Medicine',
                'store': 'orange_pharmacy',
                'formulary_source': 'local_well',  # DiaGyn = Local Well
            }
    
    return products

def parse_stock_pdf(filepath):
    """Parse Stock Summary PDF for product names and sale prices"""
    doc = fitz.open(filepath)
    products = {}
    
    for page in doc:
        text = page.get_text()
        lines = text.strip().split('\n')
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if re.match(r'^\d+$', line):
                sl_no = int(line)
                if i + 1 < len(lines):
                    item_name = lines[i + 1].strip()
                    price = None
                    if i + 2 < len(lines):
                        price_match = re.match(r'₹\s*([\d,]+\.?\d*)', lines[i + 2].strip())
                        if price_match:
                            price = float(price_match.group(1).replace(',', ''))
                    
                    if item_name and price is not None:
                        key = item_name.lower()
                        if key not in products:
                            form = infer_form(item_name)
                            orange_price = round(price * 1.15, 2)
                            
                            products[key] = {
                                'name': item_name,
                                'form': form,
                                'company': '',
                                'rx_type': '',
                                'mrp': price,
                                'sale_price': price,
                                'diagyn_price': price,
                                'orange_price': orange_price,
                                'net_rate': 0,
                                'units_per_pack': 1,
                                'composition': '',
                                'generic_name': '',
                                'category': infer_category(item_name, '', form),
                                'store': 'orange_pharmacy',
                                'formulary_source': 'vyapaar',
                            }
                    i += 3
                    continue
            i += 1
    
    return products

async def main():
    print("=" * 60)
    print("ORANGE SELECT INVENTORY IMPORT")
    print("=" * 60)
    
    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]
    
    # Parse all 3 sources
    print("\n[1/6] Parsing Orange Pharmacy XLSX (imp_new.xlsx)...")
    orange_xlsx = parse_orange_xlsx(os.path.join(DATA_DIR, 'imp_new.xlsx'))
    print(f"  Found {len(orange_xlsx)} unique products")
    
    print("\n[2/6] Parsing DiaGyn Healthcare XLSX (current_stock_new.xlsx)...")
    diagyn_xlsx = parse_diagyn_xlsx(os.path.join(DATA_DIR, 'current_stock_new.xlsx'))
    print(f"  Found {len(diagyn_xlsx)} unique products")
    
    print("\n[3/6] Parsing Stock Summary PDF (stock_summary_new.pdf)...")
    pdf_items = parse_stock_pdf(os.path.join(DATA_DIR, 'stock_summary_new.pdf'))
    print(f"  Found {len(pdf_items)} unique products")
    
    # Merge: DiaGyn takes priority over Orange, XLSX takes priority over PDF
    print("\n[4/6] Merging with pricing rules...")
    merged = {}
    
    # Start with PDF items (lowest priority - limited data)
    for key, item in pdf_items.items():
        merged[key] = item
    
    # Add Orange XLSX (overrides PDF where names match, adds new items)
    for key, item in orange_xlsx.items():
        merged[key] = item  # Full details override PDF-only entries
    
    # Add DiaGyn XLSX (highest priority - overrides everything for matching names)
    for key, item in diagyn_xlsx.items():
        if key in merged:
            # Merge: keep DiaGyn pricing but preserve any extra data from Orange
            existing = merged[key]
            item['orange_price'] = existing.get('orange_price', item['orange_price'])
            if not item.get('composition') and existing.get('composition'):
                item['composition'] = existing['composition']
            if not item.get('category') or item['category'] == 'General Medicine':
                if existing.get('category') and existing['category'] != 'General Medicine':
                    item['category'] = existing['category']
        merged[key] = item
    
    print(f"  Total merged unique products: {len(merged)}")
    
    # Prepare documents for DB
    print("\n[5/6] Preparing database documents...")
    now = datetime.now(timezone.utc).isoformat()
    documents = []
    
    for key, item in merged.items():
        doc = {
            'id': str(uuid.uuid4())[:11],
            'name': item['name'],
            'form': item['form'],
            'company': item['company'],
            'manufacturer': item['company'],
            'rx_type': item['rx_type'],
            'category': item['category'],
            'composition': item['composition'],
            'generic_name': item['generic_name'],
            'mrp': item['mrp'],
            'sale_price': item['sale_price'],
            'diagyn_price': item['diagyn_price'],
            'orange_price': item['orange_price'],
            'units_per_pack': item['units_per_pack'],
            'store': 'orange_pharmacy',
            'is_formulary': True,
            'formulary_source': item['formulary_source'],
            'formulary_tagged_at': now,
            'created_at': now,
            'in_stock': True,
        }
        documents.append(doc)
    
    # Drop existing collection and insert fresh
    print(f"\n[6/6] Inserting {len(documents)} products into trusted_formulary...")
    await db.trusted_formulary.drop()
    
    # Insert in batches
    batch_size = 200
    for i in range(0, len(documents), batch_size):
        batch = documents[i:i + batch_size]
        await db.trusted_formulary.insert_many(batch)
        print(f"  Inserted batch {i // batch_size + 1}: {len(batch)} items")
    
    # Create indexes
    await db.trusted_formulary.create_index([("name", 1)])
    await db.trusted_formulary.create_index([("category", 1)])
    await db.trusted_formulary.create_index([("formulary_source", 1)])
    await db.trusted_formulary.create_index([("name", "text")])
    
    # Verify
    total = await db.trusted_formulary.count_documents({})
    diagyn_count = await db.trusted_formulary.count_documents({"formulary_source": "local_well"})
    orange_count = await db.trusted_formulary.count_documents({"formulary_source": "vyapaar"})
    
    # Category breakdown
    pipeline = [
        {"$group": {"_id": "$category", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}}
    ]
    categories = await db.trusted_formulary.aggregate(pipeline).to_list(100)
    
    print(f"\n{'=' * 60}")
    print(f"IMPORT COMPLETE!")
    print(f"{'=' * 60}")
    print(f"Total items:     {total}")
    print(f"DiaGyn (Local Well): {diagyn_count}")
    print(f"Orange (Vyapaar):    {orange_count}")
    print(f"\nCategories:")
    for cat in categories:
        print(f"  {cat['_id']}: {cat['count']}")
    
    client.close()

if __name__ == "__main__":
    asyncio.run(main())
