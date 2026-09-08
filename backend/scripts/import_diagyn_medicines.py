"""
Import medicines from DiaGyn stock report.
Uses AI (Gemini) to categorize medicines by therapeutic system.
Only uses: Product Name, Composition, Form, Company, MRP.
"""
import asyncio
import json
import uuid
import openpyxl
from datetime import datetime, timezone
from motor.motor_asyncio import AsyncIOMotorClient
from emergentintegrations.llm.chat import LlmChat, UserMessage

MONGO_URL = "mongodb://localhost:27017"
DB_NAME = "test_database"

import os
from dotenv import load_dotenv
load_dotenv("/app/backend/.env")
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY", "")

XLSX_PATH = "/tmp/stock_products.xlsx"


def extract_products():
    wb = openpyxl.load_workbook(XLSX_PATH)
    sheet = wb.active
    products = {}
    within_dupes = []
    total_rows = 0
    for r in range(14, sheet.max_row + 1):
        name = sheet.cell(r, 1).value
        if not name or str(name).strip() == '':
            continue
        name = str(name).strip()
        total_rows += 1
        form = str(sheet.cell(r, 2).value or '').strip()
        company = str(sheet.cell(r, 3).value or '').strip()
        mrp_val = sheet.cell(r, 12).value
        composition = str(sheet.cell(r, 21).value or '').strip()
        if composition == 'None':
            composition = ''
        try:
            mrp = float(mrp_val) if mrp_val else 0.0
        except Exception:
            mrp = 0.0
        key = name.upper()
        if key in products:
            within_dupes.append((name, products[key]['name'], "within-spreadsheet duplicate (same batch)"))
        else:
            products[key] = {
                'name': name, 'form': form if form and form != 'None' else 'Other',
                'company': company if company and company != 'None' else '',
                'mrp': mrp, 'composition': composition,
            }
    return list(products.values()), within_dupes, total_rows


async def categorize_with_ai(products):
    med_list = []
    for i, p in enumerate(products):
        entry = f"{i+1}. {p['name']}"
        if p['composition']:
            entry += f" ({p['composition']})"
        med_list.append(entry)

    batch_size = 70
    all_categories = {}

    for batch_start in range(0, len(med_list), batch_size):
        batch = med_list[batch_start:batch_start + batch_size]
        batch_text = "\n".join(batch)

        prompt = f"""You are a pharmacist. Categorize each medicine into ONE therapeutic category.

Categories (pick BEST fit):
- Antibiotics
- Pain Relief & Anti-inflammatory
- Diabetes
- Cardiac & BP
- Gastro & Digestive
- Women's Health
- Respiratory & Allergy
- Vitamins & Supplements
- Skin & Dermatology
- Neuro & CNS
- Eye & ENT
- Urology & Kidney
- Bones & Joints
- Hormones & Thyroid
- Liver Care
- General & OTC

Respond ONLY with a valid JSON object mapping the number to the category.
Example: {{"1": "Antibiotics", "2": "Pain Relief & Anti-inflammatory"}}

Medicines:
{batch_text}"""

        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"med-categorize-{batch_start}",
            system_message="You are a pharmacist expert. Respond only with valid JSON."
        ).with_model("gemini", "gemini-2.5-flash")

        try:
            response = await chat.send_message(UserMessage(text=prompt))
            text = response.strip()
            if "```json" in text:
                text = text.split("```json")[1].split("```")[0].strip()
            elif "```" in text:
                text = text.split("```")[1].split("```")[0].strip()
            cats = json.loads(text)
            for k, v in cats.items():
                idx = int(k) - 1 + batch_start
                if 0 <= idx < len(products):
                    all_categories[idx] = v
            print(f"  Batch {batch_start//batch_size + 1}: categorized {len(cats)} medicines")
        except Exception as e:
            print(f"  Batch {batch_start//batch_size + 1} error: {e}")
            for i in range(batch_start, min(batch_start + batch_size, len(products))):
                if i not in all_categories:
                    all_categories[i] = "General & OTC"

    for i, p in enumerate(products):
        p['category'] = all_categories.get(i, "General & OTC")
    return products


async def main():
    products, within_dupes, total_rows = extract_products()
    print(f"Total data rows: {total_rows}")
    print(f"Unique products: {len(products)}")
    print(f"Within-sheet duplicates: {len(within_dupes)}")

    print("\nCategorizing with AI (Gemini)...")
    products = await categorize_with_ai(products)

    cat_counts = {}
    for p in products:
        cat_counts[p['category']] = cat_counts.get(p['category'], 0) + 1
    print("\nAI Category distribution:")
    for cat, count in sorted(cat_counts.items(), key=lambda x: -x[1]):
        print(f"  {cat}: {count}")

    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]
    existing = await db.medicines.find({}, {"_id": 0, "name": 1}).to_list(100000)
    existing_names = {m["name"].strip().upper() for m in existing}
    print(f"\nExisting medicines in DB: {len(existing_names)}")

    db_dupes = []
    to_insert = []
    for p in products:
        key = p['name'].strip().upper()
        if key in existing_names:
            db_dupes.append((p['name'], p['name'], "already exists in DB"))
        else:
            to_insert.append(p)

    print(f"DB duplicates (skipped): {len(db_dupes)}")
    print(f"New medicines to insert: {len(to_insert)}")

    inserted = 0
    for med in to_insert:
        doc = {
            "id": str(uuid.uuid4()),
            "name": med["name"],
            "composition": med["composition"],
            "generic_name": med["composition"],
            "form": med["form"],
            "unit": med["form"].lower(),
            "manufacturer": med["company"],
            "company": med["company"],
            "mrp": med["mrp"],
            "price": med["mrp"],
            "sale_price": med["mrp"],
            "discount_percent": 0,
            "category": med["category"],
            "description": "",
            "image_url": "",
            "prescription_required": False,
            "is_active": True,
            "store": "diagyn_healthcare",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat()
        }
        await db.medicines.insert_one(doc)
        inserted += 1

    final_count = await db.medicines.count_documents({})
    all_dupes = within_dupes + db_dupes

    print(f"\n{'='*60}")
    print(f"=== IMPORT RESULTS ===")
    print(f"Spreadsheet rows: {total_rows}")
    print(f"Within-sheet duplicates: {len(within_dupes)}")
    print(f"Already in DB (skipped): {len(db_dupes)}")
    print(f"New medicines inserted: {inserted}")
    print(f"Total medicines in DB now: {final_count}")

    print(f"\n=== ALL DUPLICATES ({len(all_dupes)}) ===")
    for name, dup_of, reason in sorted(all_dupes):
        print(f"  {name:50s} -> '{dup_of}' ({reason})")

    client.close()

if __name__ == "__main__":
    asyncio.run(main())
