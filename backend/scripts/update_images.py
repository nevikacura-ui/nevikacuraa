"""
Update image URLs for both Orange Pharmacy and Orange HealthPlus from Excel files.
Handles: new images (DB missing) + updated images (different URLs).
"""
import os, sys
import openpyxl
from pymongo import MongoClient, UpdateOne

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")

client = MongoClient(MONGO_URL)
db = client[DB_NAME]
col = db.medicines

def update_store_images(excel_path, store_name, id_col=0, name_col=1, url_col=2):
    print(f"\n=== Processing {store_name} from {excel_path} ===")
    
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    
    # Build lookup from Excel
    excel_data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[name_col]
        url = row[url_col]
        if not name:
            continue
        has_url = bool(url and str(url).strip())
        if has_url:
            excel_data[str(name).strip().upper()] = str(url).strip()
    wb.close()
    
    print(f"  Excel products with URLs: {len(excel_data):,}")
    
    # Batch update: find all products in this store
    ops = []
    new_count = 0
    updated_count = 0
    batch_size = 5000
    
    cursor = col.find(
        {"store": store_name},
        {"_id": 0, "id": 1, "name": 1, "image_url": 1}
    )
    
    for med in cursor:
        name_key = str(med.get("name", "")).strip().upper()
        if name_key not in excel_data:
            continue
        
        excel_url = excel_data[name_key]
        db_url = str(med.get("image_url", "") or "").strip()
        
        if not db_url:
            # DB missing image - add it
            ops.append(UpdateOne({"id": med["id"]}, {"$set": {"image_url": excel_url}}))
            new_count += 1
        elif excel_url.split('|')[0].strip() != db_url.split('|')[0].strip():
            # Different URL - update to new one
            ops.append(UpdateOne({"id": med["id"]}, {"$set": {"image_url": excel_url}}))
            updated_count += 1
        
        if len(ops) >= batch_size:
            col.bulk_write(ops, ordered=False)
            ops = []
    
    if ops:
        col.bulk_write(ops, ordered=False)
    
    print(f"  New images added: {new_count:,}")
    print(f"  URLs updated: {updated_count:,}")
    return new_count, updated_count

if __name__ == "__main__":
    # Process Pharmacy
    n1, u1 = update_store_images("/tmp/drugs_image_url.xlsx", "orange_pharmacy")
    
    # Process HealthPlus
    n2, u2 = update_store_images("/tmp/otc_image_url.xlsx", "orange_healthplus")
    
    print(f"\n=== TOTAL ===")
    print(f"New images: {n1 + n2:,}")
    print(f"Updated URLs: {u1 + u2:,}")
    
    # Verify
    ph_img = col.count_documents({"store": "orange_pharmacy", "image_url": {"$nin": ["", None]}})
    hp_img = col.count_documents({"store": "orange_healthplus", "image_url": {"$nin": ["", None]}})
    print(f"\nPharmacy with images now: {ph_img:,}")
    print(f"HealthPlus with images now: {hp_img:,}")
